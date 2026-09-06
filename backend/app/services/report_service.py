import csv
import io
from datetime import datetime
from fastapi import HTTPException, status
from sqlalchemy.orm import Session
from sqlalchemy import func

from app.crud.report import (
    create_report,
    get_all_reports,
    get_report_by_id,
    update_report,
    delete_report,
)

from app.models.report import Report
from app.models.vendor import Vendor
from app.models.purchase_order import PurchaseOrder
from app.models.contract import Contract
from app.models.risk import Risk
from app.models.vendor_performance import VendorPerformance

from app.schemas.report import (
    ReportCreate,
    ReportUpdate,
)


def create_report_service(
    db: Session,
    report: ReportCreate,
):
    new_report = Report(**report.model_dump())

    return create_report(
        db,
        new_report,
    )


def get_all_reports_service(
    db: Session,
):
    return get_all_reports(db)


def get_report_by_id_service(
    db: Session,
    report_id: int,
):
    report = get_report_by_id(
        db,
        report_id,
    )

    if not report:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Report not found",
        )

    return report


def update_report_service(
    db: Session,
    report_id: int,
    report_update: ReportUpdate,
):
    report = get_report_by_id(
        db,
        report_id,
    )

    if not report:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Report not found",
        )

    update_data = report_update.model_dump(
        exclude_unset=True,
    )

    for key, value in update_data.items():
        setattr(report, key, value)

    return update_report(
        db,
        report,
    )


def delete_report_service(
    db: Session,
    report_id: int,
):
    report = get_report_by_id(
        db,
        report_id,
    )

    if not report:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Report not found",
        )

    delete_report(
        db,
        report,
    )

    return {
        "message": "Report deleted successfully"
    }


def generate_and_record_report_service(
    db: Session,
    report_type: str,
    user_name: str = "Procurement User",
    file_format: str = "CSV",
):
    timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")
    report_name = f"{report_type} Report - {datetime.now().strftime('%b %d, %Y')}"

    new_report = Report(
        report_name=report_name,
        report_type=report_type,
        generated_by=user_name,
        file_format=file_format.upper(),
        status="Generated",
    )

    return create_report(db, new_report)


def export_report_csv_service(
    db: Session,
    report_id: int,
):
    report = get_report_by_id_service(db, report_id)
    return export_analytics_csv_service(db, report.report_type)


def export_analytics_csv_service(
    db: Session,
    category: str,
):
    output = io.StringIO()
    writer = csv.writer(output)
    cat_lower = category.lower()

    if "vendor" in cat_lower or "performance" in cat_lower:
        writer.writerow(["Vendor ID", "Company Name", "Category", "Status", "Email", "Phone", "Performance Score", "On-Time Deliveries", "Delayed Deliveries", "Quality Rating", "Order Completion Rate %", "Response Time (hrs)"])
        vendors = db.query(Vendor).all()
        performances = {p.vendor_id: p for p in db.query(VendorPerformance).all()}
        for v in vendors:
            p = performances.get(v.id)
            writer.writerow([
                v.id,
                v.company_name,
                v.category,
                v.status or "Active",
                v.email,
                v.phone or "",
                p.performance_score if p else "N/A",
                p.on_time_deliveries if p else 0,
                p.delayed_deliveries if p else 0,
                p.quality_rating if p else 0.0,
                p.order_completion_rate if p else 0.0,
                p.response_time if p else 0.0,
            ])
        filename = f"vrip_vendor_performance_{datetime.now().strftime('%Y%m%d')}.csv"

    elif "procurement" in cat_lower or "order" in cat_lower or "purchase" in cat_lower:
        writer.writerow(["PO Number", "Vendor ID", "Order Date", "Expected Delivery", "Total Amount ($)", "Status", "Payment Status", "Remarks"])
        pos = db.query(PurchaseOrder).all()
        for po in pos:
            writer.writerow([
                po.po_number,
                po.vendor_id,
                po.order_date,
                po.expected_delivery,
                f"{po.total_amount:.2f}",
                po.status,
                po.payment_status,
                po.remarks or "",
            ])
        filename = f"vrip_procurement_orders_{datetime.now().strftime('%Y%m%d')}.csv"

    elif "contract" in cat_lower:
        writer.writerow(["Contract Number", "Title", "Vendor ID", "Contract Value ($)", "Start Date", "End Date", "Status", "Payment Terms"])
        contracts = db.query(Contract).all()
        for c in contracts:
            writer.writerow([
                c.contract_number,
                c.contract_title or "",
                c.vendor_id,
                f"{c.contract_value:.2f}" if c.contract_value else "0.00",
                c.start_date,
                c.end_date,
                c.status,
                c.payment_terms or "",
            ])
        filename = f"vrip_contracts_{datetime.now().strftime('%Y%m%d')}.csv"

    elif "risk" in cat_lower or "compliance" in cat_lower:
        writer.writerow(["Risk ID", "Vendor ID", "Risk Type", "Severity", "Impact Score", "Status", "Description", "Created Date"])
        risks = db.query(Risk).all()
        for r in risks:
            writer.writerow([
                r.id,
                r.vendor_id,
                r.risk_type,
                r.severity,
                r.impact_score,
                r.status,
                r.description,
                r.created_at,
            ])
        filename = f"vrip_risk_compliance_{datetime.now().strftime('%Y%m%d')}.csv"

    else:
        writer.writerow(["Entity", "Count", "Total Value ($)"])
        v_count = db.query(Vendor).count()
        po_count = db.query(PurchaseOrder).count()
        po_val = db.query(PurchaseOrder).with_entities(func.coalesce(func.sum(PurchaseOrder.total_amount), 0)).scalar() or 0
        c_count = db.query(Contract).count()
        c_val = db.query(Contract).with_entities(func.coalesce(func.sum(Contract.contract_value), 0)).scalar() or 0
        r_count = db.query(Risk).count()

        writer.writerow(["Vendors", v_count, "—"])
        writer.writerow(["Purchase Orders", po_count, f"{po_val:.2f}"])
        writer.writerow(["Contracts", c_count, f"{c_val:.2f}"])
        writer.writerow(["Risk Incidents", r_count, "—"])
        filename = f"vrip_executive_summary_{datetime.now().strftime('%Y%m%d')}.csv"

    return output.getvalue(), filename


def export_report_excel_service(
    db: Session,
    report_id: int,
):
    report = get_report_by_id_service(db, report_id)
    return export_analytics_excel_service(db, report.report_type)


def export_analytics_excel_service(
    db: Session,
    category: str,
):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Procurement Audit"
    cat_lower = category.lower()

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")

    if "vendor" in cat_lower or "performance" in cat_lower:
        headers = ["Vendor ID", "Company Name", "Category", "Status", "Email", "Phone", "Performance Score", "On-Time Deliveries", "Delayed Deliveries", "Quality Rating", "Order Completion Rate %", "Response Time (hrs)"]
        ws.append(headers)
        vendors = db.query(Vendor).all()
        performances = {p.vendor_id: p for p in db.query(VendorPerformance).all()}
        for v in vendors:
            p = performances.get(v.id)
            ws.append([
                v.id,
                v.company_name,
                v.category,
                v.status or "Active",
                v.email,
                v.phone or "",
                p.performance_score if p else "N/A",
                p.on_time_deliveries if p else 0,
                p.delayed_deliveries if p else 0,
                p.quality_rating if p else 0.0,
                p.order_completion_rate if p else 0.0,
                p.response_time if p else 0.0,
            ])
        filename = f"vrip_vendor_performance_{datetime.now().strftime('%Y%m%d')}.xlsx"

    elif "procurement" in cat_lower or "order" in cat_lower or "purchase" in cat_lower:
        headers = ["PO Number", "Vendor ID", "Order Date", "Expected Delivery", "Total Amount ($)", "Status", "Payment Status", "Remarks"]
        ws.append(headers)
        pos = db.query(PurchaseOrder).all()
        for po in pos:
            ws.append([
                po.po_number,
                po.vendor_id,
                str(po.order_date),
                str(po.expected_delivery),
                float(po.total_amount),
                po.status,
                po.payment_status,
                po.remarks or "",
            ])
        filename = f"vrip_procurement_orders_{datetime.now().strftime('%Y%m%d')}.xlsx"

    elif "contract" in cat_lower:
        headers = ["Contract Number", "Title", "Vendor ID", "Contract Value ($)", "Start Date", "End Date", "Status", "Payment Terms"]
        ws.append(headers)
        contracts = db.query(Contract).all()
        for c in contracts:
            ws.append([
                c.contract_number,
                c.contract_title or "",
                c.vendor_id,
                float(c.contract_value) if c.contract_value else 0.0,
                str(c.start_date),
                str(c.end_date),
                c.status,
                c.payment_terms or "",
            ])
        filename = f"vrip_contracts_{datetime.now().strftime('%Y%m%d')}.xlsx"

    elif "risk" in cat_lower or "compliance" in cat_lower:
        headers = ["Risk ID", "Vendor ID", "Risk Type", "Severity", "Impact Score", "Status", "Description", "Created Date"]
        ws.append(headers)
        risks = db.query(Risk).all()
        for r in risks:
            ws.append([
                r.id,
                r.vendor_id,
                r.risk_type,
                r.severity,
                r.impact_score,
                r.status,
                r.description,
                str(r.created_at),
            ])
        filename = f"vrip_risk_compliance_{datetime.now().strftime('%Y%m%d')}.xlsx"

    else:
        headers = ["Entity", "Count", "Total Value ($)"]
        ws.append(headers)
        v_count = db.query(Vendor).count()
        po_count = db.query(PurchaseOrder).count()
        po_val = db.query(PurchaseOrder).with_entities(func.coalesce(func.sum(PurchaseOrder.total_amount), 0)).scalar() or 0
        c_count = db.query(Contract).count()
        c_val = db.query(Contract).with_entities(func.coalesce(func.sum(Contract.contract_value), 0)).scalar() or 0
        r_count = db.query(Risk).count()

        ws.append(["Vendors", v_count, "—"])
        ws.append(["Purchase Orders", po_count, float(po_val)])
        ws.append(["Contracts", c_count, float(c_val)])
        ws.append(["Risk Incidents", r_count, "—"])
        filename = f"vrip_executive_summary_{datetime.now().strftime('%Y%m%d')}.xlsx"

    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue(), filename


def export_report_pdf_service(
    db: Session,
    report_id: int,
):
    report = get_report_by_id_service(db, report_id)
    return export_analytics_pdf_service(db, report.report_type)


def export_analytics_pdf_service(
    db: Session,
    category: str,
):
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(letter), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'TitleStyle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        textColor=colors.HexColor('#1E3A8A'),
        spaceAfter=6,
    )
    subtitle_style = ParagraphStyle(
        'SubtitleStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        textColor=colors.HexColor('#64748B'),
        spaceAfter=14,
    )
    cell_style = ParagraphStyle(
        'CellStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        textColor=colors.HexColor('#0F172A'),
    )
    head_cell_style = ParagraphStyle(
        'HeadCellStyle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        textColor=colors.white,
    )

    story = []
    story.append(Paragraph(f"Vendor Reliability Platform (VRIP) — {category} Report", title_style))
    story.append(Paragraph(f"Generated on {datetime.now().strftime('%B %d, %Y at %H:%M:%S')} | Authoritative PostgreSQL Record", subtitle_style))
    story.append(Spacer(1, 10))

    cat_lower = category.lower()

    if "vendor" in cat_lower or "performance" in cat_lower:
        headers = ["ID", "Company", "Category", "Status", "Score", "On-Time", "Delayed", "Quality", "Fulfillment", "SLA Resp"]
        table_data = [[Paragraph(h, head_cell_style) for h in headers]]
        vendors = db.query(Vendor).all()
        performances = {p.vendor_id: p for p in db.query(VendorPerformance).all()}
        for v in vendors:
            p = performances.get(v.id)
            row = [
                str(v.id),
                v.company_name[:20],
                v.category[:15],
                v.status or "Active",
                f"{p.performance_score:.1f}" if p and p.performance_score else "N/A",
                str(p.on_time_deliveries) if p else "0",
                str(p.delayed_deliveries) if p else "0",
                f"{p.quality_rating:.1f}" if p and p.quality_rating else "N/A",
                f"{p.order_completion_rate:.0f}%" if p and p.order_completion_rate else "N/A",
                f"{p.response_time:.1f}h" if p and p.response_time else "N/A",
            ]
            table_data.append([Paragraph(cell, cell_style) for cell in row])
        filename = f"vrip_vendor_performance_{datetime.now().strftime('%Y%m%d')}.pdf"

    elif "procurement" in cat_lower or "order" in cat_lower or "purchase" in cat_lower:
        headers = ["PO Number", "Vendor ID", "Order Date", "Expected Delivery", "Total Amount", "Status", "Payment Status"]
        table_data = [[Paragraph(h, head_cell_style) for h in headers]]
        pos = db.query(PurchaseOrder).all()
        for po in pos:
            row = [
                po.po_number,
                str(po.vendor_id),
                str(po.order_date),
                str(po.expected_delivery),
                f"${po.total_amount:,.2f}",
                po.status,
                po.payment_status,
            ]
            table_data.append([Paragraph(cell, cell_style) for cell in row])
        filename = f"vrip_procurement_orders_{datetime.now().strftime('%Y%m%d')}.pdf"

    elif "contract" in cat_lower:
        headers = ["Contract #", "Title", "Vendor ID", "Value ($)", "Start Date", "End Date", "Status"]
        table_data = [[Paragraph(h, head_cell_style) for h in headers]]
        contracts = db.query(Contract).all()
        for c in contracts:
            row = [
                c.contract_number,
                (c.contract_title or "")[:25],
                str(c.vendor_id),
                f"${c.contract_value:,.2f}" if c.contract_value else "$0.00",
                str(c.start_date),
                str(c.end_date),
                c.status,
            ]
            table_data.append([Paragraph(cell, cell_style) for cell in row])
        filename = f"vrip_contracts_{datetime.now().strftime('%Y%m%d')}.pdf"

    elif "risk" in cat_lower or "compliance" in cat_lower:
        headers = ["Risk ID", "Vendor ID", "Risk Type", "Severity", "Impact", "Status", "Description"]
        table_data = [[Paragraph(h, head_cell_style) for h in headers]]
        risks = db.query(Risk).all()
        for r in risks:
            row = [
                str(r.id),
                str(r.vendor_id),
                r.risk_type,
                r.severity,
                str(r.impact_score),
                r.status,
                (r.description or "")[:40],
            ]
            table_data.append([Paragraph(cell, cell_style) for cell in row])
        filename = f"vrip_risk_compliance_{datetime.now().strftime('%Y%m%d')}.pdf"

    else:
        headers = ["Entity", "Record Count", "Total Value ($)"]
        table_data = [[Paragraph(h, head_cell_style) for h in headers]]
        v_count = db.query(Vendor).count()
        po_count = db.query(PurchaseOrder).count()
        po_val = db.query(PurchaseOrder).with_entities(func.coalesce(func.sum(PurchaseOrder.total_amount), 0)).scalar() or 0
        c_count = db.query(Contract).count()
        c_val = db.query(Contract).with_entities(func.coalesce(func.sum(Contract.contract_value), 0)).scalar() or 0
        r_count = db.query(Risk).count()

        table_data.append([Paragraph("Vendors", cell_style), Paragraph(str(v_count), cell_style), Paragraph("—", cell_style)])
        table_data.append([Paragraph("Purchase Orders", cell_style), Paragraph(str(po_count), cell_style), Paragraph(f"${po_val:,.2f}", cell_style)])
        table_data.append([Paragraph("Contracts", cell_style), Paragraph(str(c_count), cell_style), Paragraph(f"${c_val:,.2f}", cell_style)])
        table_data.append([Paragraph("Risk Incidents", cell_style), Paragraph(str(r_count), cell_style), Paragraph("—", cell_style)])
        filename = f"vrip_executive_summary_{datetime.now().strftime('%Y%m%d')}.pdf"

    t = Table(table_data, hAlign='LEFT')
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('TOPPADDING', (0, 0), (-1, 0), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('TOPPADDING', (0, 1), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 6),
    ]))
    story.append(t)

    doc.build(story)
    output.seek(0)
    return output.getvalue(), filename